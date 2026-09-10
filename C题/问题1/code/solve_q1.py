#!/usr/bin/env python3
"""Solve Question 1 with a two-stage linear program and write all deliverables."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.optimize import linprog


ROOT = Path(__file__).resolve().parents[2]
QUESTION_DIR = ROOT / "问题1"
INPUT_FILE = ROOT / "附件" / "附件1.xlsx"
TEMPLATE_FILE = ROOT / "附件" / "附件5" / "result1.xlsx"
RESULT_DIR = QUESTION_DIR / "result"
DATA_DIR = QUESTION_DIR / "data"

N = 144
DT = 1.0 / 6.0
ETA = 0.9
E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
P_MAX = 5000.0


def variable_slices() -> tuple[slice, slice, slice, slice, slice]:
    grid = slice(0, N)
    charge = slice(N, 2 * N)
    discharge = slice(2 * N, 3 * N)
    curtail = slice(3 * N, 4 * N)
    energy = slice(4 * N, 5 * N + 1)
    return grid, charge, discharge, curtail, energy


def load_input() -> tuple[list[str], np.ndarray, np.ndarray, np.ndarray]:
    workbook = load_workbook(INPUT_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if len(rows) != N:
        raise ValueError(f"附件1应包含{N}个数据点，实际为{len(rows)}个")

    source_times = [str(row[0]) for row in rows]
    price = np.array([float(row[1]) for row in rows])
    load = np.array([float(row[2]) for row in rows])
    pv = np.array([float(row[3]) for row in rows])
    if np.any(price < 0) or np.any(load < 0) or np.any(pv < 0):
        raise ValueError("附件1中的电价、负载和光伏功率不应为负数")
    return source_times, price, load, pv


def load_interval_labels() -> list[str]:
    workbook = load_workbook(TEMPLATE_FILE, read_only=True, data_only=True)
    sheet = workbook["计划购电量"]
    labels = [sheet.cell(row=i, column=1).value for i in range(2, N + 2)]
    if len(labels) != N or any(label is None for label in labels):
        raise ValueError("result1.xlsx模板中的时间段不完整")
    return [str(label) for label in labels]


def build_constraints(load: np.ndarray, pv: np.ndarray):
    grid, charge, discharge, curtail, energy = variable_slices()
    n_variables = 5 * N + 1
    a_eq = np.zeros((2 * N + 2, n_variables))
    b_eq = np.zeros(2 * N + 2)

    for t in range(N):
        # G + PV + D = L + C + W
        a_eq[t, grid.start + t] = 1.0
        a_eq[t, charge.start + t] = -1.0
        a_eq[t, discharge.start + t] = 1.0
        a_eq[t, curtail.start + t] = -1.0
        b_eq[t] = load[t] - pv[t]

        # E_(t+1) = E_t + eta*C*dt - D*dt/eta
        row = N + t
        a_eq[row, energy.start + t] = -1.0
        a_eq[row, energy.start + t + 1] = 1.0
        a_eq[row, charge.start + t] = -ETA * DT
        a_eq[row, discharge.start + t] = DT / ETA

    a_eq[2 * N, energy.start] = 1.0
    b_eq[2 * N] = E_INITIAL
    a_eq[2 * N + 1, energy.stop - 1] = 1.0
    b_eq[2 * N + 1] = E_INITIAL

    bounds = (
        [(0.0, None)] * N
        + [(0.0, P_MAX)] * N
        + [(0.0, P_MAX)] * N
        + [(0.0, None)] * N
        + [(E_MIN, E_MAX)] * (N + 1)
    )
    return a_eq, b_eq, bounds


def solve(price: np.ndarray, load: np.ndarray, pv: np.ndarray) -> np.ndarray:
    grid, charge, discharge, curtail, _ = variable_slices()
    a_eq, b_eq, bounds = build_constraints(load, pv)

    primary_objective = np.zeros(5 * N + 1)
    primary_objective[grid] = price * DT
    primary = linprog(
        primary_objective,
        A_eq=a_eq,
        b_eq=b_eq,
        bounds=bounds,
        method="highs",
    )
    if not primary.success:
        raise RuntimeError(f"第一阶段线性规划求解失败：{primary.message}")

    # Resolve cost-equivalent LP solutions without inventing a battery degradation fee.
    secondary_objective = np.zeros(5 * N + 1)
    secondary_objective[charge] = DT
    secondary_objective[discharge] = DT
    secondary_objective[curtail] = DT
    cost_tolerance = max(1e-7, abs(primary.fun) * 1e-10)
    secondary = linprog(
        secondary_objective,
        A_ub=np.array([primary_objective]),
        b_ub=np.array([primary.fun + cost_tolerance]),
        A_eq=a_eq,
        b_eq=b_eq,
        bounds=bounds,
        method="highs",
    )
    if not secondary.success:
        raise RuntimeError(f"第二阶段线性规划求解失败：{secondary.message}")
    return secondary.x


def clean(values: np.ndarray, tolerance: float = 1e-7) -> np.ndarray:
    result = values.copy()
    result[np.abs(result) < tolerance] = 0.0
    return result


def write_official_result(
    labels: list[str], grid_energy: np.ndarray, charge_energy: np.ndarray,
    discharge_energy: np.ndarray, stored_energy: np.ndarray,
) -> Path:
    output = RESULT_DIR / "result1.xlsx"
    shutil.copy2(TEMPLATE_FILE, output)
    workbook = load_workbook(output)
    purchase_sheet = workbook["计划购电量"]
    for row, (label, value) in enumerate(zip(labels, grid_energy), start=2):
        purchase_sheet.cell(row=row, column=1, value=label)
        purchase_sheet.cell(row=row, column=2, value=float(value))
        purchase_sheet.cell(row=row, column=2).number_format = "0.0000"

    storage_sheet = workbook["充放电量"]
    for block in range(6):
        start = block * 24
        stop = start + 24
        storage_sheet.cell(row=block + 2, column=2, value=float(charge_energy[start:stop].sum()))
        storage_sheet.cell(row=block + 2, column=3, value=float(discharge_energy[start:stop].sum()))
        storage_sheet.cell(row=block + 2, column=2).number_format = "0.0000"
        storage_sheet.cell(row=block + 2, column=3).number_format = "0.0000"
    storage_sheet.cell(row=2, column=5, value=float(stored_energy[0]))
    storage_sheet.cell(row=3, column=5, value=float(stored_energy[-1]))
    storage_sheet.cell(row=2, column=5).number_format = "0.0000"
    storage_sheet.cell(row=3, column=5).number_format = "0.0000"
    workbook.save(output)
    return output


def write_detail_files(
    source_times: list[str], labels: list[str], price: np.ndarray,
    load: np.ndarray, pv: np.ndarray, grid_power: np.ndarray,
    charge_power: np.ndarray, discharge_power: np.ndarray,
    curtail_power: np.ndarray, stored_energy: np.ndarray,
) -> None:
    headers = [
        "序号", "模板时间段", "附件时间", "电价(元/kWh)", "负载功率(kW)",
        "光伏功率(kW)", "购电功率(kW)", "购电量(kWh)", "充电功率(kW)",
        "充电量(kWh)", "放电功率(kW)", "放电量(kWh)", "弃光功率(kW)",
        "弃光量(kWh)", "时段初储电量(kWh)", "时段末储电量(kWh)", "购电费(元)",
    ]
    records = []
    for t in range(N):
        records.append([
            t + 1, labels[t], source_times[t], price[t], load[t], pv[t],
            grid_power[t], grid_power[t] * DT, charge_power[t], charge_power[t] * DT,
            discharge_power[t], discharge_power[t] * DT, curtail_power[t],
            curtail_power[t] * DT, stored_energy[t], stored_energy[t + 1],
            price[t] * grid_power[t] * DT,
        ])

    csv_path = DATA_DIR / "问题1完整计划.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(records)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "完整计划"
    sheet.append(headers)
    for record in records:
        sheet.append([float(value) if isinstance(value, np.floating) else value for value in record])
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [8, 20, 14, 16, 17, 17, 17, 17, 17, 17, 17, 17, 17, 17, 21, 21, 15]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2, min_col=4):
        for cell in row:
            cell.number_format = "0.0000"
    workbook.save(DATA_DIR / "问题1完整计划.xlsx")


def write_markdown(
    labels: list[str], price: np.ndarray, load: np.ndarray, pv: np.ndarray,
    grid_power: np.ndarray, charge_power: np.ndarray, discharge_power: np.ndarray,
    curtail_power: np.ndarray, stored_energy: np.ndarray,
) -> None:
    grid_energy = grid_power * DT
    charge_energy = charge_power * DT
    discharge_energy = discharge_power * DT
    curtail_energy = curtail_power * DT
    cost = float(np.dot(price, grid_energy))
    baseline_grid = np.maximum(load - pv, 0.0) * DT
    baseline_cost = float(np.dot(price, baseline_grid))
    selected = ["10:00-10:10", "12:00-12:10", "16:00-16:10", "18:00-18:10"]
    index = {label: i for i, label in enumerate(labels)}

    table1_rows = "\n".join(
        f"| {label} | {grid_energy[index[label]]:.4f} |" for label in selected
    )
    table2_rows = []
    for block in range(6):
        start = block * 24
        stop = start + 24
        table2_rows.append(
            f"| {block * 4}:00-{(block + 1) * 4}:00 | "
            f"{charge_energy[start:stop].sum():.4f} | {discharge_energy[start:stop].sum():.4f} |"
        )

    simultaneous = int(np.count_nonzero((charge_power > 1e-6) & (discharge_power > 1e-6)))
    max_balance_error = float(np.max(np.abs(grid_power + pv + discharge_power - load - charge_power - curtail_power)))
    state_error = stored_energy[1:] - stored_energy[:-1] - ETA * charge_power * DT + discharge_power * DT / ETA
    max_state_error = float(np.max(np.abs(state_error)))

    content = f"""# C题问题一：微网日前计划购电策略

## 1. 问题分析

根据附件1，全天共有144个数据点，依次与官方 `result1.xlsx` 模板中的144个10分钟区间对应。令时段长度 $\\Delta t=1/6\\ \mathrm{{h}}$。在每天0:00已知全天电价、小区负载和光伏预测功率的条件下，需要联合决定每个时段的外网购电、储能充电、储能放电和弃光功率，使负载始终得到满足，并使24:00储电量恢复至初始值。

该问题的目标函数和全部约束均为线性形式，因此采用线性规划求解。第一阶段最小化购电费；考虑到线性规划可能存在多个等价最优解，第二阶段在购电费保持最优的前提下最小化充电量、放电量与弃光量之和，用于排除无意义的能量循环。第二阶段不改变题目规定的经济目标。按照已确认的数据口径，附件1中从 `0:10` 至 `0:00+1` 的144个数据点依次映射到模板的144行，不另行平移或插值。

## 2. 模型假设

1. 附件1给出的负载、光伏预测功率和电价在对应10分钟时段内保持不变。
2. 题目给出的“充放电效率为90%”解释为充电和放电的单程效率均为0.9。
3. 微网可以弃光，但不向外网售电；弃光不产生收益与费用。
4. 不考虑储能自放电、设备启停成本和寿命折损，储能仅受题目给出的容量、功率和效率限制。
5. 初始储电量采用附录1给出的6000 kWh，且24:00恢复为6000 kWh。

## 3. 符号说明

| 符号 | 含义 | 单位 |
|---|---|---|
| $L_t$ | 第$t$时段小区负载功率 | kW |
| $P_t^{{\mathrm{{PV}}}}$ | 第$t$时段光伏预测功率 | kW |
| $p_t$ | 第$t$时段外网电价 | 元/kWh |
| $G_t$ | 第$t$时段购电功率 | kW |
| $C_t$ | 第$t$时段充电功率 | kW |
| $D_t$ | 第$t$时段放电功率 | kW |
| $W_t$ | 第$t$时段弃光功率 | kW |
| $E_t$ | 第$t$时段开始时的储电量 | kWh |

## 4. 线性规划模型

### 4.1 目标函数

全天购电费最小：

$$
\\min Z=\\sum_{{t=1}}^{{144}}p_tG_t\\Delta t.
$$

### 4.2 功率平衡

每个时段均满足：

$$
G_t+P_t^{{\\mathrm{{PV}}}}+D_t=L_t+C_t+W_t.
$$

### 4.3 储能状态转移

取单程效率 $\\eta=0.9$，则：

$$
E_{{t+1}}=E_t+\\eta C_t\\Delta t-\\frac{{D_t}}{{\\eta}}\\Delta t.
$$

### 4.4 运行边界与首尾约束

$$
1200\\le E_t\\le10800,
$$

$$
0\\le C_t\\le5000,\\qquad 0\\le D_t\\le5000,
$$

$$
G_t\\ge0,\\qquad W_t\\ge0,
$$

$$
E_0=E_{{144}}=6000.
$$

### 4.5 等价最优解筛选

记第一阶段得到的最小购电费为 $Z^*$。在约束 $Z\\le Z^*+\\varepsilon$ 下求解辅助目标：

$$
\\min \\sum_{{t=1}}^{{144}}(C_t+D_t+W_t)\\Delta t,
$$

其中 $\\varepsilon$ 仅取求解器数值容差量级。该步骤用于从购电费相同的策略中选取能量周转和弃光更少的方案，不向原题增加经济成本参数。

## 5. 求解方法

使用 Python 的 `scipy.optimize.linprog` 与 HiGHS 求解器计算。程序首先读取附件1的144行数据，并按照官方模板的行序逐项对应；随后建立包含购电、充电、放电、弃光及145个储电状态变量的线性规划；最后填写官方结果模板并输出完整明细。

## 6. 计算结果

### 6.1 指定时段购电量及全天指标

| 时间段 | 购电量（kWh） |
|---|---:|
{table1_rows}
| **全天购电量** | **{grid_energy.sum():.4f}** |
| **全天购电费** | **{cost:.4f} 元** |

### 6.2 分时段充放电量及首尾储电量

| 时间段 | 充电量（kWh） | 放电量（kWh） |
|---|---:|---:|
{chr(10).join(table2_rows)}

| 时刻 | 储电量（kWh） |
|---|---:|
| 0:00 | {stored_energy[0]:.4f} |
| 24:00 | {stored_energy[-1]:.4f} |

全天累计充电量为 **{charge_energy.sum():.4f} kWh**，累计放电量为 **{discharge_energy.sum():.4f} kWh**，弃光量为 **{curtail_energy.sum():.4f} kWh**。储能运行期间最低储电量为 **{stored_energy.min():.4f} kWh**，最高储电量为 **{stored_energy.max():.4f} kWh**。

若不使用储能，仅由光伏优先供负载、缺额由外网补足，则基准购电费为 {baseline_cost:.4f} 元。优化后购电费减少 **{baseline_cost - cost:.4f} 元**，降幅为 **{(baseline_cost - cost) / baseline_cost * 100:.4f}%**。

## 7. 可行性校验

| 校验项 | 结果 |
|---|---:|
| 最大功率平衡误差 | {max_balance_error:.3e} kW |
| 最大储能状态转移误差 | {max_state_error:.3e} kWh |
| 同时充放电时段数 | {simultaneous} |
| 期末与期初储电量偏差 | {abs(stored_energy[-1] - stored_energy[0]):.3e} kWh |

计算结果满足功率平衡、储能容量、充放电功率和首尾储电量约束。

## 8. 输出文件

- `result/result1.xlsx`：按附件5官方模板填写的提交文件。
- `data/问题1完整计划.xlsx`：含输入数据、各决策变量、逐时段储电量和费用的完整明细。
- `data/问题1完整计划.csv`：便于后续程序读取和复核的明细数据。
- `code/solve_q1.py`：问题一完整求解和结果生成程序。
"""
    (QUESTION_DIR / "问题1.md").write_text(content, encoding="utf-8")


def verify(
    grid: np.ndarray, charge: np.ndarray, discharge: np.ndarray,
    curtail: np.ndarray, energy: np.ndarray, load: np.ndarray, pv: np.ndarray,
) -> None:
    tolerance = 1e-5
    balance_error = grid + pv + discharge - load - charge - curtail
    state_error = energy[1:] - energy[:-1] - ETA * charge * DT + discharge * DT / ETA
    checks = {
        "功率平衡": np.max(np.abs(balance_error)) <= tolerance,
        "储能递推": np.max(np.abs(state_error)) <= tolerance,
        "储电量范围": energy.min() >= E_MIN - tolerance and energy.max() <= E_MAX + tolerance,
        "充放电功率": charge.max() <= P_MAX + tolerance and discharge.max() <= P_MAX + tolerance,
        "首尾储电量": abs(energy[0] - E_INITIAL) <= tolerance and abs(energy[-1] - E_INITIAL) <= tolerance,
        "非负变量": min(grid.min(), charge.min(), discharge.min(), curtail.min()) >= -tolerance,
        "无同时充放电": not np.any((charge > tolerance) & (discharge > tolerance)),
    }
    failed = [name for name, passed in checks.items() if not passed]
    if failed:
        raise RuntimeError("结果校验失败：" + "、".join(failed))


def main() -> None:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    source_times, price, load, pv = load_input()
    labels = load_interval_labels()
    solution = solve(price, load, pv)
    grid_slice, charge_slice, discharge_slice, curtail_slice, energy_slice = variable_slices()
    grid = clean(solution[grid_slice])
    charge = clean(solution[charge_slice])
    discharge = clean(solution[discharge_slice])
    curtail = clean(solution[curtail_slice])
    energy = clean(solution[energy_slice])

    verify(grid, charge, discharge, curtail, energy, load, pv)
    write_official_result(labels, grid * DT, charge * DT, discharge * DT, energy)
    write_detail_files(
        source_times, labels, price, load, pv, grid, charge, discharge, curtail, energy,
    )
    write_markdown(labels, price, load, pv, grid, charge, discharge, curtail, energy)

    total_cost = float(np.dot(price, grid * DT))
    print(f"求解成功：全天购电量 {grid.sum() * DT:.4f} kWh")
    print(f"全天购电费：{total_cost:.4f} 元")
    print(f"累计充/放/弃光：{charge.sum() * DT:.4f} / {discharge.sum() * DT:.4f} / {curtail.sum() * DT:.4f} kWh")
    print(f"SOC范围：{energy.min():.4f} - {energy.max():.4f} kWh")


if __name__ == "__main__":
    main()
