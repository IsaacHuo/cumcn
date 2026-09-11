"""Read-only forecast diagnostics for Q2/Q3 (no solver or result mutation)."""
from pathlib import Path
import csv, json, math
import numpy as np
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data'/'raw'; OUT=ROOT/'results'; TABLE=ROOT/'paper'/'tables'
OUT.mkdir(exist_ok=True); TABLE.mkdir(exist_ok=True)

def metrics(err):
    e=np.asarray(err,float); e=e[np.isfinite(e)]
    return {'bias':float(e.mean()),'mae':float(np.abs(e).mean()),'rmse':float(np.sqrt(np.mean(e*e))),'n':int(e.size)}
def read_q2():
    files=list(DATA.glob('附件2.xlsx'))
    if not files: raise FileNotFoundError('data/raw/附件2.xlsx')
    wb=load_workbook(files[0],read_only=True,data_only=True)
    out={}
    for name in wb.sheetnames:
        rows=list(wb[name].values); dates=[r[0].date().isoformat() for r in rows[1:]]
        out[name]=np.asarray([r[1:] for r in rows[1:]],float)/6
    wb.close(); return dates,out
def q2_diag():
    dates,sheets=read_q2(); load=sheets["小区负载"]; pv=sheets["光伏发电实际功率"]
    start=dates.index('2025-02-01'); end=dates.index('2025-12-31')+1
    errs={'load':[],'pv':[],'net_load':[]}
    rows=[]
    for d in range(start,end):
        if d<7: continue
        fl=np.r_[load[d-7],load[d-6]][:144]; fp=np.tile(pv[d-1],2)[:144]
        for key,e in [('load',fl-load[d]),('pv',fp-pv[d]),('net_load',(fl-fp)-(load[d]-pv[d]))]: errs[key].extend(e)
        rows.append({'date':dates[d],'load':metrics(fl-load[d]),'pv':metrics(fp-pv[d]),'net_load':metrics((fl-fp)-(load[d]-pv[d]))})
    return {'period':['2025-02-01','2025-12-31'],'q2':{k:metrics(v) for k,v in errs.items()},'q2_daily':rows}
def read_q3():
    f=list(DATA.glob('附件3.xlsx'))[0]; wb=load_workbook(f,read_only=True,data_only=True); ws=wb.active
    heads=[ws.cell(1,j).value for j in range(1,27)]; rows=[]
    current=None
    for row in ws.iter_rows(min_row=2,values_only=True):
        date=row[0]
        if date:
            if hasattr(date,'date'): date=date.date().isoformat()
            else: date=str(date)
        if date: current=date
        rows.append((current,row[1],np.asarray(row[2:26],float)))
    wb.close(); return rows
def q3_diag():
    dates,sheets=read_q2(); pv=sheets['光伏发电实际功率']; q3=read_q3(); idx={d:i for i,d in enumerate(dates)}
    by_release={}; by_lead={}
    for i,(date,t,fc) in enumerate(q3):
        if not date or date<'2025-02-01': continue
        di=idx.get(date); hour=int(str(t).split(':')[0]);
        if di is None: continue
        if di+(hour+24)//24 >= len(pv): continue
        # Attachment 3 forecasts are kW; attachment 2 is stored as kWh/10 min.
        actual=np.array([pv[di+(hour+j)//24,(hour+j)%24*6:(hour+j)%24*6+6].mean()*6 for j in range(1,25)])
        e=fc-actual
        key=str(t)
        by_release.setdefault(key,[]).extend(e)
        for a,b in [(1,6),(7,12),(13,18),(19,24)]: by_lead.setdefault(f'{a}-{b}h',[]).extend(e[a-1:b])
    return {'period':['2025-02-01','2025-12-31'],'q3_by_release':{k:metrics(v) for k,v in by_release.items()},'q3_by_lead':{k:metrics(v) for k,v in by_lead.items()}}
def main():
    result={'q2':q2_diag(),'q3':q3_diag()}
    (OUT/'forecast_diagnostics.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    with (OUT/'forecast_diagnostics.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f); w.writerow(['section','group','bias','mae','rmse','n'])
        for sec,items in [('q2',result['q2']['q2']),('q3_release',result['q3']['q3_by_release']),('q3_lead',result['q3']['q3_by_lead'])]:
            for k,v in items.items(): w.writerow([sec,k,v['bias'],v['mae'],v['rmse'],v['n']])
    tex=['\\begin{tabular}{llrrrr}','\\toprule','部分 & 分组 & Bias & MAE & RMSE & $n$ \\\\','\\midrule']
    for sec,items in [('Q2',result['q2']['q2']),('Q3发布时刻',result['q3']['q3_by_release']),('Q3提前量',result['q3']['q3_by_lead'])]:
        for k,v in items.items(): tex.append(f'{sec} & {k} & {v["bias"]:.2f} & {v["mae"]:.2f} & {v["rmse"]:.2f} & {v["n"]} \\\\')
    tex += ['\\bottomrule','\\end{tabular}','']
    (TABLE/'forecast_diagnostics.tex').write_text('\n'.join(tex),encoding='utf-8')
if __name__=='__main__': main()
