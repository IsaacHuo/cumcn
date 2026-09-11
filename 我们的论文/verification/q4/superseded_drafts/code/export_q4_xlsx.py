"""Export Q4 workbooks; use --allow-partial only with pilot directories."""
import argparse,shutil
from datetime import datetime
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
def unmerge(ws):
 for r in list(ws.merged_cells.ranges):ws.unmerge_cells(str(r))
def lab(i):return f'{i//6:02d}:{(i%6)*10:02d}-{(i+1)//6:02d}:{((i+1)%6)*10:02d}' if i<143 else '23:50-24:00'
def export(v,t,out,partial=False):
 folder=ROOT/'results/q4'/v
 if partial and not folder.exists(): folder=ROOT/'results/q4'/('pilot_'+v)
 fs=sorted(folder.glob('2025-*.npz'));assert len(fs)==334 or partial
 if not fs:return
 shutil.copyfile(t,out);wb=load_workbook(out);[unmerge(s) for s in wb.worksheets]; plans=[wb.worksheets[0]]+([wb.worksheets[1]] if len(wb.worksheets)==4 else []); storage=wb.worksheets[-2]; emergency=wb.worksheets[-1]
 for s in plans:
  for j,h in enumerate(['日期\\时间']+[lab(i) for i in range(144)]+(['全天调整购电量','全天总费用（含紧急购电）'] if s is plans[-1] and len(plans)==2 else ['全天计划购电量','全天总费用（计划）']),1):s.cell(1,j).value=h
 for r,f in enumerate(fs,2):
  with np.load(f) as z:
   dt=datetime.strptime(f.stem,'%Y-%m-%d');vals=[z['adjusted_kwh'] if len(plans)==2 else z['plan_kwh']]
   for s,arr,key in zip(plans,vals,['total_cost' if len(plans)==2 else 'plan_cost']):s.cell(r,1).value=dt;[s.cell(r,i+2).__setattr__('value',float(x)) for i,x in enumerate(arr)];s.cell(r,146).value=float(arr.sum());s.cell(r,147).value=float(z[key])
   if len(plans)==2:
    s=plans[0];arr=z['plan_kwh'];s.cell(r,1).value=dt;[s.cell(r,i+2).__setattr__('value',float(x)) for i,x in enumerate(arr)];s.cell(r,146).value=float(arr.sum());s.cell(r,147).value=float(z['plan_cost'])
 storage.delete_rows(2,max(0,storage.max_row-1));emergency.delete_rows(2,max(0,emergency.max_row-1));rr=2;er=2
 for f in fs:
  with np.load(f) as z:
   dt=datetime.strptime(f.stem,'%Y-%m-%d')
   for i in range(6):storage.cell(rr,1).value=dt if i==0 else None;storage.cell(rr,2).value=f'{i*4:02d}:00-{(i+1)*4:02d}:00';storage.cell(rr,3).value=float(z['charge_kwh'][i*24:i*24+24].sum());storage.cell(rr,4).value=float(z['discharge_kwh'][i*24:i*24+24].sum());storage.cell(rr,5).value='0:00' if i==0 else '24:00' if i==1 else None;storage.cell(rr,6).value=float(z['soc_kwh'][0] if i==0 else z['soc_kwh'][-1]) if i<2 else None;rr+=1
   x=z['emergency_kwh'];i=0;events=[]
   while i<144:
    if x[i]<=1e-5:i+=1;continue
    j=i+1
    while j<144 and x[j]>1e-5:j+=1
    events.append((lab(i)[:6]+'-'+lab(j-1)[7:],float(x[i:j].sum())));i=j
   if not events:events=[('无',0.)]
   for k,(it,a) in enumerate(events):emergency.cell(er,1).value=dt if k==0 else None;emergency.cell(er,2).value=it;emergency.cell(er,3).value=a;er+=1
 wb.save(out)
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--allow-partial',action='store_true');a=p.parse_args();export('q2',ROOT/'data/raw/result4-2_template.xlsx',ROOT/'results/result4-2.xlsx',a.allow_partial);export('main',ROOT/'data/raw/result4-3_template.xlsx',ROOT/'results/result4-3.xlsx',a.allow_partial)
