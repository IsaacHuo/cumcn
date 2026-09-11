"""Causal Q2 stochastic look-ahead control. Units: kWh, yuan.

Future scenario recourse is a planning relaxation. Only the common current
action is executed; subsequent realized data never enter a controller call.
"""
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.deps'))
import argparse,csv,datetime,hashlib,json,time
import numpy as np
from scipy.sparse import coo_matrix
from openpyxl import load_workbook
import highspy

N=144;H=288;ETA=.9;CAP=5000/6;LOW=1200.;HIGH=10800.;INITIAL=6000.
TOL=1e-5
PRICE_FILE=ROOT/'data/raw/附件1.xlsx'
DATA_FILE=ROOT/'data/raw/附件2.xlsx'

def read_inputs():
    wb=load_workbook(PRICE_FILE,read_only=True,data_only=True)
    price=np.array([r[1] for r in list(wb.active.values)[1:]],float);wb.close()
    wb=load_workbook(DATA_FILE,read_only=True,data_only=True)
    arrays=[];dates=None
    for name in ['小区负载','光伏发电实际功率']:
        rows=list(wb[name].values);ds=[r[0].date().isoformat() for r in rows[1:]]
        if dates is None: dates=ds
        assert ds==dates
        for i,v in enumerate(rows[0][1:],1):
            minute=1440 if v=='0:00+1' else (v.hour*60+v.minute if hasattr(v,'hour') else int(v.split(':')[0])*60+int(v.split(':')[1]))
            assert minute==10*i
        arrays.append(np.asarray([r[1:] for r in rows[1:]],float)/6)
    wb.close()
    assert arrays[0].shape==(365,144) and arrays[1].shape==(365,144)
    assert all(np.isfinite(a).all() and (a>=0).all() for a in arrays)
    assert len(set(dates))==365 and (price>0).all()
    return price,arrays[0],arrays[1],dates

def forecast(history_load,history_pv):
    assert len(history_load)>=7
    return np.r_[history_load[-7],history_load[-6]],np.tile(history_pv[-1],2)

def scenarios(history_load,history_pv,count=7):
    """Input is strictly the completed historical prefix, never target-day data."""
    day=len(history_load); fl,fp=forecast(history_load,history_pv)
    origins=np.arange(max(7,day-29),day-1)
    el=np.array([np.r_[history_load[i]-history_load[i-7],history_load[i+1]-history_load[i-6]] for i in origins])
    ep=np.array([np.r_[history_pv[i]-history_pv[i-1],history_pv[i+1]-history_pv[i-1]] for i in origins])
    assert len(origins)>0 and (origins+1<day).all()
    scale_l=np.maximum(el.std(axis=0),1/6);scale_p=np.maximum(ep.std(axis=0),1/6)
    features=np.concatenate([el/scale_l,ep/scale_p],axis=1)
    distance=np.mean((features[:,None,:]-features[None,:,:])**2,axis=2)
    k=min(count,len(origins));net_error=(el-ep).sum(axis=1)
    if k==1: medoids=[int(np.argmin(distance.sum(axis=1)))]
    else:
        medoids=list(dict.fromkeys([int(np.argmin(net_error)),int(np.argmax(net_error))]))
        while len(medoids)<k:
            ds=distance[:,medoids].min(axis=1);ds[medoids]=-1
            medoids.append(int(np.argmax(ds)))
        # Keep both extreme observations, refine other medoids deterministically.
        for _ in range(10):
            assign=np.argmin(distance[:,medoids],axis=1);new=list(medoids)
            for j in range(2,k):
                members=np.flatnonzero(assign==j)
                if len(members):new[j]=int(members[np.argmin(distance[np.ix_(members,members)].sum(axis=1))])
            if new==medoids:break
            medoids=new
    medoids=np.array(medoids,int)
    assign=np.argmin(distance[:,medoids],axis=1)
    weights=np.bincount(assign,minlength=k).astype(float)/len(origins)
    keep=weights>0;medoids=medoids[keep];weights=weights[keep]
    return dict(load=np.maximum(fl+el[medoids],0),pv=np.maximum(fp+ep[medoids],0),prior=weights,
                forecast_load=fl,forecast_pv=fp,error_load=el[medoids],error_pv=ep[medoids],
                scale_load=scale_l,scale_pv=scale_p,origins=origins[medoids],pool_origins=origins)

class HorizonLP:
    """Reuse a HiGHS basis while advancing measured RHS/bounds within one day."""
    def __init__(self,price,loads,pvs,prob,soc,target=6000.,capacity=CAP):
        self.price=np.tile(price,2);self.loads=np.array(loads);self.pvs=np.array(pvs)
        self.S=len(loads);self.B=5*H+1;self.size=H+self.S*self.B;self.capacity=capacity
        self.t=-1;self.target=target
        rr=[];cc=[];vv=[];lo=[];up=[]
        def add(cols,vals,l,u):
            r=len(lo);rr.extend([r]*len(cols));cc.extend(cols);vv.extend(vals);lo.append(l);up.append(u);return r
        self.balance=np.empty((self.S,H),int);self.state=np.empty_like(self.balance);self.spill=np.empty_like(self.balance)
        lower=np.zeros(self.size);upper=np.full(self.size,np.inf)
        for s in range(self.S):
            b=self.base(s);upper[b:b+2*H]=capacity
            lower[b+4*H:b+5*H+1]=LOW;upper[b+4*H:b+5*H+1]=HIGH
            lower[b+4*H]=upper[b+4*H]=soc
            lower[b+5*H]=upper[b+5*H]=target
            for t in range(H):
                rhs=loads[s,t]-pvs[s,t]
                self.balance[s,t]=add([t,b+t,b+H+t,b+2*H+t,b+3*H+t],[1,-1,1,1,-1],rhs,rhs)
                self.state[s,t]=add([b+4*H+t,b+4*H+t+1,b+t,b+H+t],[-1,1,-ETA,1/ETA],0,0)
                self.spill[s,t]=add([b+3*H+t,t],[1,-1],-np.inf,pvs[s,t])
        self.links=[]
        for s in range(1,self.S):
            for kind in [0,1]:
                self.links.append((add([self.base(s)+kind*H,self.base(0)+kind*H],[1,-1],-np.inf,np.inf),s,kind))
        a=coo_matrix((vv,(rr,cc)),shape=(len(lo),self.size)).tocsc()
        lp=highspy.HighsLp();lp.num_col_=self.size;lp.num_row_=len(lo)
        lp.col_cost_=self.objective(prob);lp.col_lower_=lower;lp.col_upper_=upper
        lp.row_lower_=np.asarray(lo);lp.row_upper_=np.asarray(up)
        lp.a_matrix_.format_=highspy.MatrixFormat.kColwise;lp.a_matrix_.start_=a.indptr.astype(np.int32)
        lp.a_matrix_.index_=a.indices.astype(np.int32);lp.a_matrix_.value_=a.data
        self.h=highspy.Highs()
        for name,value in [('output_flag',False),('threads',1),('solver','simplex'),('random_seed',2026),('primal_feasibility_tolerance',1e-8),('dual_feasibility_tolerance',1e-8),('time_limit',60.)]:self.h.setOptionValue(name,value)
        self.h.passModel(lp);self.last_status='not_run';self.recoveries=[]

    def base(self,s):return H+s*self.B

    def objective(self,prob):
        c=np.zeros(self.size);c[:H]=self.price
        for s,w in enumerate(prob):c[self.base(s)+2*H:self.base(s)+3*H]=w*5*self.price
        return c

    def solve(self):
        self.h.run();status=self.h.getModelStatus();self.last_status=str(status)
        if status!=highspy.HighsModelStatus.kOptimal:
            original=str(status)
            # A numerically stale warm basis can return kUnknown. Re-solve the
            # exact same LP from a fresh basis, never change constraints/data.
            self.h.clearSolver();self.h.run();status=self.h.getModelStatus()
            self.last_status=str(status)
            self.recoveries.append({'interval':self.t,'initial_status':original,'retry_status':str(status)})
        if status!=highspy.HighsModelStatus.kOptimal:raise RuntimeError(self.last_status)
        return np.asarray(self.h.getSolution().col_value)

    def fix_plan(self,plan):
        ids=np.arange(N,dtype=np.int32);self.h.changeColsBounds(N,ids,np.asarray(plan),np.asarray(plan))
        self.plan=np.asarray(plan)

    def step(self,t,soc,load,pv,prob):
        assert t==self.t+1
        columns=[];cl=[];cu=[];rowids=[];rl=[];ru=[]
        def bound(i,l,u):columns.append(i);cl.append(l);cu.append(u)
        def rb(i,l,u):rowids.append(i);rl.append(l);ru.append(u)
        net=load-pv-self.plan[t]
        for s in range(self.S):
            b=self.base(s)
            if t>0:
                for kind in range(4):bound(b+kind*H+t-1,0,0)
                for table in [self.balance,self.state,self.spill]:rb(table[s,t-1],-np.inf,np.inf)
            bound(b+4*H+t,soc,soc)
            rb(self.balance[s,t],load-pv,load-pv)
            rb(self.spill[s,t],-np.inf,pv)
            if net>=0:
                bound(b+t,0,0);bound(b+H+t,0,min(self.capacity,net,max(0,(soc-LOW)*ETA)))
                bound(b+2*H+t,0,net);bound(b+3*H+t,0,0)
            else:
                bound(b+t,0,min(self.capacity,-net,max(0,(HIGH-soc)/ETA)));bound(b+H+t,0,0)
                bound(b+2*H+t,0,0);bound(b+3*H+t,0,-net)
        for row,s,kind in self.links:
            if t>0:
                self.h.changeCoeff(row,self.base(s)+kind*H+t-1,0)
                self.h.changeCoeff(row,self.base(0)+kind*H+t-1,0)
            self.h.changeCoeff(row,self.base(s)+kind*H+t,1)
            self.h.changeCoeff(row,self.base(0)+kind*H+t,-1)
            rb(row,0,0)
        self.h.changeColsBounds(len(columns),np.asarray(columns,np.int32),np.asarray(cl,float),np.asarray(cu,float))
        self.h.changeRowsBounds(len(rowids),np.asarray(rowids,np.int32),np.asarray(rl,float),np.asarray(ru,float))
        cost=self.objective(prob)
        ids=np.concatenate([np.arange(self.base(s)+2*H,self.base(s)+3*H,dtype=np.int32) for s in range(self.S)])
        self.h.changeColsCost(len(ids),ids,cost[ids]);self.t=t
        x=self.solve();b=self.base(0)
        c,d,e,w=[float(x[b+k*H+t]) for k in range(4)]
        assert all(abs(float(x[self.base(s)+k*H+t])-x[b+k*H+t])<TOL for s in range(self.S) for k in range(4))
        return c,d,e,w

def fallback(soc,net,capacity=CAP):
    if net>=0:
        d=min(net,capacity,max(0,(soc-LOW)*ETA));return 0.,d,net-d,0.
    c=min(-net,capacity,max(0,(HIGH-soc)/ETA));return c,0.,0.,-net-c

def posterior(sc,observed_load,observed_pv):
    n=len(observed_load)
    residual_l=np.asarray(observed_load)-sc['forecast_load'][:n]
    residual_p=np.asarray(observed_pv)-sc['forecast_pv'][:n]
    distance=.5*np.mean(((residual_l-sc['error_load'][:,:n])/sc['scale_load'][:n])**2+((residual_p-sc['error_pv'][:,:n])/sc['scale_pv'][:n])**2,axis=1)
    logw=np.log(sc['prior'])-.5*distance;w=np.exp(logw-logw.max());w/=w.sum()
    return .95*w+.05*sc['prior']

def weighted_quantile(values,weights,q=.8):
    order=np.argsort(values);return float(values[order[np.searchsorted(np.cumsum(weights[order]),q,side='left')]])

def simulate_day(price,history_load,history_pv,measurements,soc,variant='main',count=7,target=6000.):
    sc=scenarios(history_load,history_pv,count);prior=sc['prior'];S=len(prior)
    start=time.perf_counter()
    if variant=='no_storage':
        net=sc['load']-sc['pv']
        plan=np.array([max(0,weighted_quantile(net[:,t],prior)) for t in range(N)])
        expected_em=float(np.sum(prior[:,None]*5*price*np.maximum(net[:,:N]-plan,0)))
        controller=None
    else:
        if variant=='deterministic':
            planner=HorizonLP(price,sc['forecast_load'][None,:],sc['forecast_pv'][None,:],np.ones(1),soc,target)
        else:planner=HorizonLP(price,sc['load'],sc['pv'],prior,soc,target)
        x=planner.solve();plan=np.maximum(x[:N],0)
        expected_em=sum(float(np.dot(5*price,x[planner.base(s)+2*H:planner.base(s)+2*H+N]))*w for s,w in enumerate(np.ones(1) if variant=='deterministic' else prior))
        controller=planner if variant=='main' else HorizonLP(price,sc['load'],sc['pv'],prior,soc,target)
        controller.fix_plan(plan)
    plan_seconds=time.perf_counter()-start
    result=dict(plan_kwh=plan,forecast_load_kwh=sc['forecast_load'][:N],forecast_pv_kwh=sc['forecast_pv'][:N],
                scenario_load=sc['load'],scenario_pv=sc['pv'],scenario_prior=prior,scenario_origins=sc['origins'],pool_origins=sc['pool_origins'],
                scenario_error_load=sc['error_load'],scenario_error_pv=sc['error_pv'],scenario_scale_load=sc['scale_load'],scenario_scale_pv=sc['scale_pv'],
                expected_emergency_cost=expected_em,plan_seconds=plan_seconds)
    records=[];soc_values=[soc];weights=[];failures=[];ol=[];op=[];solve_seconds=[]
    for t,(load,pv) in enumerate(measurements):
        assert t<N
        ol.append(float(load));op.append(float(pv));prob=posterior(sc,ol,op);weights.append(prob)
        tick=time.perf_counter();net=load-pv-plan[t]
        if variant=='no_storage':c,d,e,w=0.,0.,max(net,0),max(-net,0)
        else:
            try:c,d,e,w=controller.step(t,soc,load,pv,prob)
            except RuntimeError as exc:
                failures.append({'interval':t,'status':str(exc)});c,d,e,w=fallback(soc,net)
        solve_seconds.append(time.perf_counter()-tick)
        # Reconcile only numerical round-off, without modifying the optimized action.
        c=max(0,c);d=max(0,d);e=max(0,e);w=max(0,w)
        nextsoc=soc+ETA*c-d/ETA
        assert abs(plan[t]+pv+d+e-load-c-w)<TOL
        assert LOW-TOL<=nextsoc<=HIGH+TOL and c<=CAP+TOL and d<=CAP+TOL
        assert not(c>TOL and (d>TOL or e>TOL))
        records.append([c,d,e,w]);soc=float(np.clip(nextsoc,LOW,HIGH));soc_values.append(soc)
    assert len(records)==N
    a=np.array(records)
    for k,name in enumerate(['charge_kwh','discharge_kwh','emergency_kwh','unused_kwh']):result[name]=a[:,k]
    result.update(soc_kwh=np.array(soc_values),load_kwh=np.array(ol),pv_kwh=np.array(op),weights=np.array(weights),solve_seconds=np.array(solve_seconds),
                  plan_cost=float(np.dot(price,plan)),emergency_cost=float(np.dot(5*price,a[:,2])),failures_json=json.dumps(failures),elapsed_seconds=time.perf_counter()-start)
    result['total_cost']=result['plan_cost']+result['emergency_cost']
    result['recoveries_json']=json.dumps(controller.recoveries if controller is not None else [])
    return result

def config(variant,count,target):
    return dict(version=1,variant=variant,scenario_count=count,target=target,history_pairs=28,prior_mix=.05,
                highspy=highspy.Highs().version(),solver_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                input_sha256=hashlib.sha256(DATA_FILE.read_bytes()).hexdigest())

def run(variant='main',count=7,target=6000.,end=365,tag=None):
    price,L,P,dates=read_inputs();tag=tag or variant
    folder=ROOT/'results/q2'/tag;folder.mkdir(parents=True,exist_ok=True)
    cfg=config(variant,count,target);cfgpath=folder/'config.json'
    if cfgpath.exists():
        previous=json.loads(cfgpath.read_text(encoding='utf-8'))
        if previous!=cfg:raise RuntimeError(f'Checkpoint configuration changed: {tag}. Use a new --tag.')
    else:cfgpath.write_text(json.dumps(cfg,indent=2),encoding='utf-8')
    soc=INITIAL
    for day in range(31,end):
        file=folder/f'{dates[day]}.npz'
        if file.exists():
            with np.load(file) as r:
                assert abs(float(r['soc_kwh'][0])-soc)<TOL
                soc=float(r['soc_kwh'][-1])
            continue
        result=simulate_day(price,L[:day],P[:day],zip(L[day],P[day]),soc,variant,count,target)
        result['date']=dates[day];result['day_index']=day
        temporary=file.with_suffix('.tmp.npz');np.savez_compressed(temporary,**result);temporary.replace(file)
        soc=float(result['soc_kwh'][-1])
        print(json.dumps({'variant':tag,'date':dates[day],'cost':round(result['total_cost'],2),'emergency_kwh':round(float(result['emergency_kwh'].sum()),2),
                          'soc_end':round(soc,2),'seconds':round(result['elapsed_seconds'],2),'failures':len(json.loads(result['failures_json']))}),flush=True)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--variant',choices=['main','deterministic','no_storage'],default='main')
    parser.add_argument('--count',type=int,default=7);parser.add_argument('--target',type=float,default=6000.);parser.add_argument('--end',type=int,default=365);parser.add_argument('--tag')
    args=parser.parse_args();run(args.variant,args.count,args.target,args.end,args.tag)
