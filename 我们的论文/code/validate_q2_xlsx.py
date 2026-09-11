"""Independent full-workbook checks for the question-2 XLSX export."""
import json
import math
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "results" / "result2.xlsx"
SOURCE = ROOT / "results" / "q2_export.json"
OUT = ROOT / "verification" / "q2_xlsx_independent_checks.json"


def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def serial_date(value):
    return value.date() if isinstance(value, datetime) else value


def main():
    src = json.loads(SOURCE.read_text(encoding="utf-8"))
    days = src["days"]
    assert isinstance(days, list) and len(days) == 334
    wb = load_workbook(RESULT, read_only=True, data_only=True)
    expected_sheets = ["计划购电量", "充放电量", "紧急购电量"]
    assert wb.sheetnames == expected_sheets, wb.sheetnames
    plan, storage, emergency = wb.worksheets
    # Artifact-tool exports may omit worksheet dimension metadata; count rows
    # through the read-only iterators rather than trusting max_row.
    plan_rows = list(plan.iter_rows())
    storage_rows = list(storage.iter_rows())
    emergency_rows = list(emergency.iter_rows())
    assert len(plan_rows) == 335 and len(plan_rows[0]) == 147
    assert len(storage_rows) == 2005 and len(storage_rows[0]) == 6
    expected_emergency_rows = sum(max(1, len(d.get("emergency_intervals") or [])) for d in days) + 1
    assert len(emergency_rows) == expected_emergency_rows, (len(emergency_rows), expected_emergency_rows)
    assert plan_rows[0][0].value == "日期\\时间"
    assert plan_rows[0][145].value == "全天计划购电量"
    assert plan_rows[0][146].value == "全天总费用（计划+紧急）"
    def interval_label(index):
        start, end = index * 10, index * 10 + 10
        return f"{start // 60:02d}:{start % 60:02d}-{end // 60:02d}:{end % 60:02d}"
    assert [c.value for c in plan_rows[0][1:145]] == [interval_label(i) for i in range(144)]
    expected_dates = [datetime.strptime(d["date"], "%Y-%m-%d").date() for d in days]
    assert expected_dates == [expected_dates[0] + timedelta(days=i) for i in range(334)]
    for i, day in enumerate(days, start=2):
        row_values = plan_rows[i - 1]
        assert serial_date(row_values[0].value) == expected_dates[i - 2]
        vals = [c.value for c in row_values[1:145]]
        assert all(finite(v) for v in vals)
        expected = day["plan_kwh"]
        assert isinstance(expected, list) and len(expected) == 144
        assert all(finite(v) for v in expected)
        expected = [float(v) for v in expected]
        assert all(abs(v - e) < 1e-7 for v, e in zip(vals, expected)), i
        assert finite(day["plan_total"]) and finite(day["total_cost"])
        assert abs(row_values[145].value - sum(vals)) < 1e-7
        assert abs(row_values[145].value - float(day["plan_total"])) < 1e-7
        assert abs(row_values[146].value - float(day["total_cost"])) < 1e-7
    for di, day in enumerate(days):
        base = 2 + di * 6
        blocks = day.get("blocks")
        assert len(blocks) == 6
        for j in range(6):
            row = base + j
            row_values = storage_rows[row - 1]
            assert row_values[1].value
            assert finite(row_values[2].value) and finite(row_values[3].value)
            block = blocks[j]
            assert block.get("interval") == row_values[1].value
            assert finite(block.get("charge_kwh")) and finite(block.get("discharge_kwh"))
            assert abs(row_values[2].value - float(block.get("charge_kwh", 0))) < 1e-7
            assert abs(row_values[3].value - float(block.get("discharge_kwh", 0))) < 1e-7
        first, second = storage_rows[base - 1], storage_rows[base]
        assert serial_date(first[0].value) == expected_dates[di]
        assert first[4].value == "0:00"
        assert finite(day.get("soc_start")) and finite(day.get("soc_end"))
        assert abs(first[5].value - float(day["soc_start"])) < 1e-7
        assert second[4].value == "24:00"
        assert abs(second[5].value - float(day["soc_end"])) < 1e-7
    row = 2
    for day, date in zip(days, expected_dates):
        events = day.get("emergency_intervals") or []
        if not events:
            row_values = emergency_rows[row - 1]
            assert serial_date(row_values[0].value) == date
            assert row_values[1].value == "无" and row_values[2].value == 0
            row += 1
            continue
        for j, event in enumerate(events):
            row_values = emergency_rows[row - 1]
            assert serial_date(row_values[0].value) == (date if j == 0 else None)
            assert row_values[1].value == str(event.get("interval", "无"))
            assert finite(event.get("kwh"))
            assert abs(row_values[2].value - float(event["kwh"])) < 1e-7
            row += 1
    residual = []
    for ws in (plan, storage, emergency):
        for cells in ws.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str) and "⁝" in cell.value:
                    residual.append(f"{ws.title}!{cell.coordinate}")
    assert not residual, residual
    result = {"sheets": wb.sheetnames, "plan_days": len(days), "plan_intervals": 144,
        "storage_rows": len(storage_rows) - 1, "emergency_rows": len(emergency_rows) - 1,
              "date_continuity": True, "template_residuals": residual,
              "plan_total_checks": len(days), "storage_boundary_checks": len(days),
              "emergency_day_checks": len(days)}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"VALIDATION_FAILED: {exc}", file=sys.stderr)
        raise
