"""Independent read-only validation for the question-3 XLSX export."""
import json, math, sys
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "results" / "result3.xlsx"
SOURCE = ROOT / "results" / "q3_export.json"
OUT = ROOT / "verification" / "q3导出检查" / "q3_xlsx_independent_checks.json"

def finite(value): return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)
def serial_date(value): return value.date() if isinstance(value, datetime) else value
def interval_label(i):
    start, end = i * 10, i * 10 + 10
    return f"{start // 60:02d}:{start % 60:02d}-{end // 60:02d}:{end % 60:02d}"
def close(a, b): return finite(a) and finite(b) and abs(float(a) - float(b)) < 1e-7

def main():
    src = json.loads(SOURCE.read_text(encoding="utf-8")); days = src.get("days")
    assert isinstance(days, list) and len(days) == 334
    wb = load_workbook(RESULT, read_only=True, data_only=True)
    expected_sheets = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    assert wb.sheetnames == expected_sheets, wb.sheetnames
    plan, adjusted, storage, emergency = wb.worksheets
    plan_rows, adjusted_rows = list(plan.iter_rows()), list(adjusted.iter_rows())
    storage_rows, emergency_rows = list(storage.iter_rows()), list(emergency.iter_rows())
    assert len(plan_rows) == len(adjusted_rows) == 335 and len(plan_rows[0]) == len(adjusted_rows[0]) == 147
    assert len(storage_rows) == 2005 and len(storage_rows[0]) == 6
    expected_emergency_rows = sum(max(1, len(d.get("emergency_intervals") or [])) for d in days) + 1
    assert len(emergency_rows) == expected_emergency_rows
    for ws in (plan, adjusted): assert [c.value for c in ws[1][1:145]] == [interval_label(i) for i in range(144)]
    assert plan_rows[0][0].value == adjusted_rows[0][0].value == "日期\\时间"
    assert plan_rows[0][145].value == "全天计划购电量" and plan_rows[0][146].value == "全天总费用（计划）"
    assert adjusted_rows[0][145].value == "全天调整购电量" and adjusted_rows[0][146].value == "全天总费用（含紧急购电）"
    expected_dates = [datetime.strptime(d["date"], "%Y-%m-%d").date() for d in days]
    assert expected_dates == [expected_dates[0] + timedelta(days=i) for i in range(334)]
    for i, day in enumerate(days):
        for rows, key, total_key, cost_key in [(plan_rows, "plan_kwh", "plan_total", "plan_cost"), (adjusted_rows, "adjusted_kwh", "adjusted_total", "total_cost")]:
            row = rows[i + 1]; assert serial_date(row[0].value) == expected_dates[i]
            vals = [c.value for c in row[1:145]]; expected = day.get(key)
            assert isinstance(expected, list) and len(expected) == 144 and all(finite(v) for v in vals + expected)
            assert all(close(v, e) for v, e in zip(vals, expected))
            assert close(row[145].value, sum(vals)) and close(row[145].value, day.get(total_key)) and close(row[146].value, day.get(cost_key))
        blocks = day.get("blocks"); assert isinstance(blocks, list) and len(blocks) == 6
        base = 1 + i * 6
        for j, block in enumerate(blocks):
            row = storage_rows[base + j]; assert row[1].value == block.get("interval") and finite(row[2].value) and finite(row[3].value)
            assert close(row[2].value, block.get("charge_kwh")) and close(row[3].value, block.get("discharge_kwh"))
        first, second = storage_rows[base], storage_rows[base + 1]
        assert serial_date(first[0].value) == expected_dates[i] and first[4].value == "0:00" and close(first[5].value, day.get("soc_start"))
        assert second[4].value == "24:00" and close(second[5].value, day.get("soc_end"))
    row = 1
    for day, date in zip(days, expected_dates):
        events = day.get("emergency_intervals") or []
        if not events:
            vals = emergency_rows[row]; assert serial_date(vals[0].value) == date and vals[1].value == "无" and vals[2].value == 0; row += 1; continue
        for j, event in enumerate(events):
            vals = emergency_rows[row]; assert serial_date(vals[0].value) == (date if j == 0 else None) and vals[1].value == event.get("interval") and close(vals[2].value, event.get("kwh")); row += 1
    residual = [f"{ws.title}!{cell.coordinate}" for ws in wb.worksheets for row_cells in ws.iter_rows() for cell in row_cells if isinstance(cell.value, str) and "⁝" in cell.value]
    assert not residual, residual
    result = {"sheets": wb.sheetnames, "plan_days": len(days), "plan_intervals": 144, "storage_rows": len(storage_rows) - 1, "emergency_rows": len(emergency_rows) - 1, "date_continuity": True, "template_residuals": residual, "plan_total_checks": len(days), "adjusted_total_checks": len(days), "storage_boundary_checks": len(days), "emergency_day_checks": len(days)}
    OUT.parent.mkdir(parents=True, exist_ok=True); OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"); print(json.dumps(result, ensure_ascii=False))
if __name__ == "__main__":
    try: main()
    except Exception as exc: print(f"VALIDATION_FAILED: {exc}", file=sys.stderr); raise
