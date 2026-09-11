"""Independent scalar recomputation from exported CSV, without optimizer matrices."""
from pathlib import Path
import csv,json,math
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
TOL=1e-5

def main():
    with (ROOT/'results/q1_detail.csv').open(encoding='utf-8-sig') as f: rows=list(csv.DictReader(f))
    assert len(rows)==144
    data=json.loads((ROOT/'results/q1.json').read_text(encoding='utf-8'))
    wb=load_workbook(ROOT/'data/raw/附件1.xlsx',read_only=True,data_only=True)
    raw=list(wb.active.values)[1:];wb.close()
    residuals=[];states=[];bounds=[];cost=0;energy=6000;simultaneous=0
    for i,r in enumerate(rows):
        x={k:float(v) for k,v in r.items() if k not in ['interval','source_end']}
        assert all(math.isfinite(v) for v in x.values())
        expected=f'{i//6:02d}:{i%6*10:02d}-{(i+1)//6:02d}:{(i+1)%6*10:02d}'
        assert r['interval']==expected
        assert abs(x['price_yuan_per_kwh']-raw[i][1])<1e-12
        assert abs(x['load_kwh']-raw[i][2]/6)<1e-10
        assert abs(x['pv_kwh']-raw[i][3]/6)<1e-10
        g,c,d,w=[x[k+'_kwh'] for k in ['grid','charge','discharge','curtail']]
        residuals.append(abs(g+x['pv_kwh']+d-x['load_kwh']-c-w))
        states.append(abs(energy-x['soc_start_kwh']))
        energy+=.9*c-d/.9
        states.append(abs(energy-x['soc_end_kwh']))
        bounds += [max(0,-min(g,c,d,w)),max(0,c-5000/6,d-5000/6,w-x['pv_kwh'],1200-energy,energy-10800)]
        cost+=g*x['price_yuan_per_kwh']
        assert abs(x['cost_yuan']-g*x['price_yuan_per_kwh'])<TOL
        simultaneous+=int(c>TOL and d>TOL)
    s=data['summary']
    group_error=0
    for i,b in enumerate(data['blocks']):
        for k in ['charge_kwh','discharge_kwh']:
            group_error=max(group_error,abs(sum(float(r[k]) for r in rows[i*24:(i+1)*24])-b[k]))
    report={'interval_count':len(rows),'max_balance_error_kwh':max(residuals),'max_reconstructed_soc_error_kwh':max(states),'max_bound_violation_kwh':max(bounds),'terminal_error_kwh':abs(energy-6000),'cost_error_yuan':abs(cost-s['cost_yuan']),'block_error_kwh':group_error,'simultaneous_periods':simultaneous,'whole_day_energy_error_kwh':abs(s['grid_kwh']+s['pv_kwh']-s['load_kwh']-s['curtail_kwh']-s['loss_kwh']),'tolerance':TOL}
    assert simultaneous==0
    assert all(v<TOL for k,v in report.items() if 'error' in k or 'violation' in k)
    for k in ['grid_kwh','charge_kwh','discharge_kwh','curtail_kwh','cost_yuan']:
        assert abs(sum(float(r[k]) for r in rows)-s[k])<TOL
    report['passed']=True
    (ROOT/'verification/q1_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,indent=2))

    alt=json.loads((ROOT/'results/q1_efficiency_sensitivity.json').read_text(encoding='utf-8'))
    eta=alt['eta']; energy=6000.; errors=[]; alt_cost=0.
    for i,r in enumerate(alt['rows']):
        g,c,d,w=[r[k+'_kwh'] for k in ['grid','charge','discharge','curtail']]
        errors.extend([abs(g+raw[i][3]/6+d-raw[i][2]/6-c-w),abs(energy-r['soc_start_kwh'])])
        energy+=eta*c-d/eta
        errors.extend([abs(energy-r['soc_end_kwh']),max(0,-g,-c,-d,-w,c-5000/6,d-5000/6,w-raw[i][3]/6,1200-energy,energy-10800)])
        assert not(c>TOL and d>TOL)
        alt_cost+=raw[i][1]*g
    errors.extend([abs(energy-6000),abs(alt_cost-alt['summary']['cost_yuan'])])
    assert max(errors)<TOL
    (ROOT/'verification/efficiency_checks.json').write_text(json.dumps({'eta':eta,'max_error':max(errors),'passed':True},indent=2),encoding='utf-8')

    # Optional final export check: run again after Excel export.
    result=ROOT/'results/result1.xlsx'
    if result.exists():
        w=load_workbook(result,read_only=False,data_only=True)
        assert w.sheetnames==['计划购电量','充放电量']
        assert w.worksheets[0].max_row==145 and w.worksheets[0].max_column==2
        assert w.worksheets[1].max_row==7 and w.worksheets[1].max_column==5
        for i,r in enumerate(data['rows'],2):
            assert w.worksheets[0].cell(i,1).value==r['interval']
            assert abs(w.worksheets[0].cell(i,2).value-r['grid_kwh'])<TOL
        for i,b in enumerate(data['blocks'],2):
            assert abs(w.worksheets[1].cell(i,2).value-b['charge_kwh'])<TOL
            assert abs(w.worksheets[1].cell(i,3).value-b['discharge_kwh'])<TOL
        assert w.worksheets[1]['E2'].value==6000 and w.worksheets[1]['E3'].value==6000
        w.close()
        (ROOT/'verification/xlsx_checks.json').write_text(json.dumps({'all_144_intervals_and_values_match':True,'all_six_blocks_match':True,'template_sheet_structure_preserved':True,'passed':True},indent=2),encoding='utf-8')

if __name__=='__main__':main()
