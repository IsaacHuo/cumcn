"""Reopen XLSX results and reconcile every output with unrounded records."""
from solve_q4 import *
from openpyxl import load_workbook

def main():
    results=[]
    for variant,name in [('q2','result4-2'),('main','result4-3')]:
        source=json.loads((ROOT/f'results/q4/{variant}_export.json').read_text(encoding='utf-8'))['days']
        wb=load_workbook(ROOT/f'results/{name}.xlsx',read_only=True,data_only=True)
        plan=wb.worksheets[0];adjusted=wb.worksheets[1] if variant=='main' else None
        storage=wb.worksheets[-2];emergency=wb.worksheets[-1]
        for sheet,field,total,cost in [(plan,'plan_kwh','plan_total','total_cost' if variant=='q2' else 'plan_cost')]+([(adjusted,'adjusted_kwh','adjusted_total','total_cost')] if adjusted else []):
            values=list(sheet.values);assert len(values)==335 and len(values[0])==147
            assert values[0][1]=='00:00-00:10' and values[0][144]=='23:50-24:00'
            for row,d in zip(values[1:],source):
                assert row[0].date().isoformat()==d['date']
                assert max(abs(np.asarray(row[1:145],float)-d[field]))<1e-8
                assert abs(float(row[145])-d[total])<TOL and abs(float(row[146])-d[cost])<TOL
        values=list(storage.values);assert len(values)==2005
        for k,d in enumerate(source):
            for i,b in enumerate(d['blocks']):
                row=values[1+k*6+i];assert row[1]==b['interval']
                assert abs(row[2]-b['charge_kwh'])<TOL and abs(row[3]-b['discharge_kwh'])<TOL
                if i==0:assert row[0].date().isoformat()==d['date'] and abs(row[5]-d['soc_start'])<TOL
                if i==1:assert abs(row[5]-d['soc_end'])<TOL
        observed={};date=None
        for row in list(emergency.values)[1:]:
            if row[0]:date=row[0].date().isoformat()
            observed.setdefault(date,[]).append((row[1],float(row[2])))
        assert len(observed)==334
        for d in source:
            expected=[(e['interval'],e['kwh']) for e in d['emergency_intervals']] or [('无',0.)]
            assert len(expected)==len(observed[d['date']])
            for a,b in zip(expected,observed[d['date']]):assert a[0]==b[0] and abs(a[1]-b[1])<TOL
            assert abs(sum(a[1] for a in observed[d['date']])-d['emergency_total'])<.0015
        results.append(dict(workbook=name,days=334,plan_intervals=48096,storage_rows=2004,worksheets=wb.sheetnames,passed=True));wb.close()
    (ROOT/'verification/q4/workbooks/validation.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(results,ensure_ascii=False))

if __name__=='__main__':main()
